from flask import Flask, render_template, request, redirect, session
from config import users, notes, tasks, files,history,events
from bson.objectid import ObjectId
from flask import flash
import os
import re 
from werkzeug.utils import secure_filename
from flask import send_from_directory,send_file
from datetime import datetime
from openai import OpenAI
from werkzeug.security import generate_password_hash, check_password_hash
import json
from pymongo import MongoClient
from google import genai
from config import GEMINI_API_KEY, GROQ_API_KEY
import zipfile
import uuid
import PIL.Image
import traceback
from flask import request, session, redirect, render_template, jsonify
from bson.objectid import ObjectId
import shutil
import stat
from datetime import datetime
from werkzeug.utils import secure_filename
from groq import Groq
import pandas as pd
import fitz
import pytesseract
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
from PIL import Image
import csv
import io
import base64
from pathlib import Path


ai_client = genai.Client(api_key=GEMINI_API_KEY)
groq_client = Groq(api_key=GROQ_API_KEY)

app = Flask(__name__)
app.secret_key = "zyntra_secret_key"

UPLOAD_FOLDER = "uploads"
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

# ============================================================
# ZYNTRA FILE CREATOR CONFIG
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

GENERATED_DIR = BASE_DIR / "generated_files"

GENERATED_DIR.mkdir(
    parents=True,
    exist_ok=True
)



# MongoDB Client
client = MongoClient(os.getenv("MONGODB_URI"))
db = client["Zyntra"]

users = db["users"]

def save_history(username, analysis_type, title, result,provider=""):

    history.insert_one({

        "username": username,
        "type": analysis_type,
        "provider": provider,
        "title": title,
        "result": result,
        "date": datetime.now()

    })

@app.route("/history")
def history_page():

    if "fullname" not in session:
        return redirect("/login")

    history_data = history.find({

        "$or": [
            {"username": session["fullname"]},
            {"fullname": session["fullname"]}
        ]

    }).sort("date", -1)

    return render_template(
        "history.html",
        history=history_data
    )

@app.route("/")
def home():
    return render_template("index.html")

@app.route("/login", methods=["GET", "POST"])
def login():

    if request.method == "POST":

        email = request.form["email"]
        password = request.form["password"]

        user = users.find_one({
            "email": email
        })

        if user and check_password_hash(user["password"], password):

            session["fullname"] = user["fullname"]
            session["email"] = user["email"]

            return redirect("/dashboard")

        return render_template(
            "login.html",
            error="Invalid Email or Password"
        )

    return render_template("login.html")

@app.route("/forgot-password", methods=["GET", "POST"])
def forgot_password():

    if request.method == "POST":

        email = request.form["email"]
        new_password = request.form["new_password"]
        confirm_password = request.form["confirm_password"]

       

        user = users.find_one({
            "email": email
        })

        if not user:
            flash("Email does not exist.", "danger")
            return render_template("forgot_password.html")

        if new_password != confirm_password:
            flash("Passwords do not match.", "danger")
            return render_template("forgot_password.html")

        hashed_password = generate_password_hash(new_password)


        users.update_one(
            {"email": email},
            {
                "$set": {
                    "password": hashed_password
                }
            }
        )

        flash("Password updated successfully. Please login.", "success")
        
    return render_template("forgot_password.html")


@app.route("/signup", methods=["GET", "POST"])
def signup():

    if request.method == "POST":

        fullname = request.form["fullname"]
        email = request.form["email"]
        password = request.form["password"]
        confirm = request.form["confirmPassword"]

        if password != confirm:

            return "Passwords do not match"

        user = users.find_one({
            "email": email
        })

        if user:

            return "Email already exists"

        users.insert_one({

            "fullname": fullname,
            "email": email,
            "password": password,

            "phone": "",
            "bio": "",
            "profile_image": "",

            "created_at": datetime.now()

        })

        return redirect("/login")

    return render_template("signup.html")


@app.route("/dashboard")
def dashboard():

    if "fullname" not in session:
        return redirect("/login")

    total_notes = notes.count_documents({
        "fullname": session["fullname"]
    })

    total_tasks = tasks.count_documents({
        "fullname": session["fullname"]
    })

    completed_tasks = tasks.count_documents({
        "fullname": session["fullname"],
        "status":"Completed"
    })

    pending_tasks = tasks.count_documents({
        "fullname": session["fullname"],
        "status": "Pending"
    })

    in_progress_tasks = tasks.count_documents({
        "fullname": session["fullname"],
        "status": "In Progress"
    })

    total_files = files.count_documents({
    "fullname": session["fullname"]
    })

    total_ai = history.count_documents({
    "fullname": session["fullname"]
    })

    recent_tasks = tasks.find({
        "fullname": session["fullname"]
    }).sort("_id", -1).limit(5)

    recent_files = files.find({
        "fullname": session["fullname"]
    }).sort("_id", -1).limit(5)

    total_events = events.count_documents({
        "fullname": session["fullname"]
    })

    recent_notes = list(notes.find({
        "fullname": session["fullname"]
    }).sort("_id", -1).limit(5))

    now = datetime.utcnow()

    for note in recent_notes:

        if "created_at" in note:

            diff = now - note["created_at"]
            seconds = diff.total_seconds()

            if seconds < 60:
                note["time_text"] = "Just now"

            elif seconds < 3600:
                mins = int(seconds // 60)
                note["time_text"] = f"{mins} min ago"

            elif seconds < 86400:
                hrs = int(seconds // 3600)
                note["time_text"] = f"{hrs} hr ago"
   
            elif seconds < 172800:
                note["time_text"] = "Yesterday"

            else:
                days = diff.days
                note["time_text"] = f"{days} days ago"

        else:
            note["time_text"] = ""

 
    return render_template(
        "dashboard.html",
        
        fullname=session["fullname"],
        total_events=total_events,
        total_notes=total_notes,
        total_tasks=total_tasks,
        completed_tasks=completed_tasks,
        pending_tasks=pending_tasks,
        total_files=total_files,
        total_ai=total_ai,
        recent_notes=recent_notes,
        recent_tasks=recent_tasks,
        recent_files=recent_files,
        in_progress_tasks=in_progress_tasks
  
    )

@app.context_processor
def inject_topbar_data():

    if "fullname" not in session:

        return {
            "current_user": None,
            "notification_count": 0,
            "pending_tasks": 0,
            "total_events": 0,
            "total_notes": 0
        }


    fullname = session["fullname"]


    current_user = users.find_one({
        "fullname": fullname
    })


    pending_tasks = tasks.count_documents({
        "fullname": fullname,
        "status": "Pending"
    })


    total_events = events.count_documents({
        "fullname": fullname
    })


    total_notes = notes.count_documents({
        "fullname": fullname
    })


    notification_count = (
        pending_tasks
        + total_events
        + total_notes
    )


    return {

        "current_user": current_user,

        "pending_tasks": pending_tasks,

        "total_events": total_events,

        "total_notes": total_notes,

        "notification_count": notification_count

    }

@app.route("/calendar")
def calendar():

    if "fullname" not in session:
        return redirect("/login")

    all_events = events.find({
        "fullname": session["fullname"]
    }).sort("date", 1)

    return render_template(
        "calendar.html",
        fullname=session["fullname"],
        events=all_events
    )

@app.route("/add-event", methods=["POST"])
def add_event():

    if "fullname" not in session:
        return redirect("/login")

    events.insert_one({

        "fullname": session["fullname"],

        "title": request.form["title"],

        "description": request.form["description"],

        "date": request.form["date"],

        "time": request.form["time"]

    })

    return redirect("/calendar")

@app.route("/edit-event/<event_id>", methods=["GET", "POST"])
def edit_event(event_id):

    if "fullname" not in session:
        return redirect("/login")

    event = events.find_one({
        "_id": ObjectId(event_id),
        "fullname": session["fullname"]
    })

    if request.method == "POST":

        events.update_one(
            {"_id": ObjectId(event_id)},
            {
                "$set": {
                    "title": request.form["title"],
                    "description": request.form["description"],
                    "date": request.form["date"],
                    "time": request.form["time"]
                }
            }
        )

        return redirect("/calendar")

    return render_template(
        "edit_event.html",
        event=event
    )


@app.route("/delete-event/<event_id>")
def delete_event(event_id):

    if "fullname" not in session:
        return redirect("/login")

    events.delete_one({
        "_id": ObjectId(event_id),
        "fullname": session["fullname"]
    })

    return redirect("/calendar")


@app.route("/logout")
def logout():

    session.clear()

    return redirect("/login")

@app.route("/notes", methods=["GET", "POST"])
def notes_page():

    if "fullname" not in session:
        return redirect("/login")

    # Add Note
    if request.method == "POST":

        title = request.form["title"]
        description = request.form["description"]

        notes.insert_one({
            "fullname": session["fullname"],
            "title": title,
            "description": description,
            "created_at": datetime.utcnow()
        })

        return redirect("/notes")

    # Search
    search = request.args.get("search")

    if search:

        all_notes = list(notes.find({
            "fullname": session["fullname"],
            "$or": [
                {
                    "title": {
                        "$regex": search,
                        "$options": "i"
                    }
                },
                {
                    "description": {
                        "$regex": search,
                        "$options": "i"
                    }
                }
            ]
        }))

    else:

        all_notes = list(notes.find({
            "fullname": session["fullname"]
        }))

    # Time Ago
    now = datetime.utcnow()

    for note in all_notes:

        if "created_at" in note:

            diff = now - note["created_at"]
            seconds = diff.total_seconds()

            if seconds < 60:
                note["time_text"] = "Just now"

            elif seconds < 3600:
                mins = int(seconds // 60)
                note["time_text"] = f"{mins} min ago"

            elif seconds < 86400:
                hrs = int(seconds // 3600)
                note["time_text"] = f"{hrs} hr ago"

            elif seconds < 172800:
                note["time_text"] = "Yesterday"

            else:
                days = diff.days
                note["time_text"] = f"{days} days ago"

        else:
            note["time_text"] = ""

    return render_template(
        "notes.html",
        notes=all_notes,
        fullname=session["fullname"]
    )

@app.route("/tasks", methods=["GET", "POST"])
def tasks_page():

    if "fullname" not in session:
        return redirect("/login")

    if request.method == "POST":

        task = request.form["task"]

        tasks.insert_one({
            "fullname": session["fullname"],
            "task": task,
            "status": "Pending"
        })

        return redirect("/tasks")

    search = request.args.get("search")

    query = {
        "fullname": session["fullname"]
    }

    if search:
        query["task"] = {
            "$regex": search,
            "$options": "i"
        }

    all_tasks = tasks.find(query)

    total_tasks = tasks.count_documents({
        "fullname": session["fullname"]
    })

    completed_tasks = tasks.count_documents({
        "fullname": session["fullname"],
        "status": "Completed"
    })

    in_progress_tasks = tasks.count_documents({
        "fullname": session["fullname"],
        "status": "In Progress"
    })

    pending_tasks = tasks.count_documents({
        "fullname": session["fullname"],
        "status": "Pending"
    })

    return render_template(
        "tasks.html",
        tasks=all_tasks,
        fullname=session["fullname"],
        total_tasks=total_tasks,
        completed_tasks=completed_tasks,
        in_progress_tasks=in_progress_tasks,
        pending_tasks=pending_tasks
    )
@app.route("/delete-task/<task_id>")
def delete_task(task_id):

    if "fullname" not in session:
        return redirect("/login")

    tasks.delete_one({
        "_id": ObjectId(task_id),
        "fullname": session["fullname"]
    })

    return redirect("/tasks")

@app.route("/edit-task/<task_id>", methods=["GET", "POST"])
def edit_task(task_id):

    if "fullname" not in session:
        return redirect("/login")

    task = tasks.find_one({
        "_id": ObjectId(task_id),
        "fullname": session["fullname"]
    })

    if request.method == "POST":

        new_task = request.form["task"]
        new_status = request.form["status"]

        tasks.update_one(
            {"_id": ObjectId(task_id)},
            {
                "$set": {
                    "task": new_task,
                    "status": new_status
                }
            }
        )

        return redirect("/tasks")

    return render_template(
        "edit_task.html",
        task=task
    )

@app.route("/delete-note/<note_id>")
def delete_note(note_id):

    if "fullname" not in session:
        return redirect("/login")

    notes.delete_one({
        "_id": ObjectId(note_id),
        "fullname": session["fullname"]
    })

    return redirect("/notes")

@app.route("/edit-note/<note_id>", methods=["GET", "POST"])
def edit_note(note_id):

    if "fullname" not in session:
        return redirect("/login")

    note = notes.find_one({
        "_id": ObjectId(note_id)
    })

    if request.method == "POST":

        title = request.form["title"]
        description = request.form["description"]

        notes.update_one(
            {"_id": ObjectId(note_id)},
            {
                "$set": {
                    "title": title,
                    "description": description
                }
            }
        )

        return redirect("/notes")

    return render_template(
        "edit_note.html",
        note=note
    )

@app.route("/files", methods=["GET", "POST"])
def files_page():

    if "fullname" not in session:
        return redirect("/login")

    if request.method == "POST":

        uploaded_file = request.files["file"]

        if uploaded_file.filename != "":

            filename = secure_filename(uploaded_file.filename)

            uploaded_file.save(
                os.path.join(
                    app.config["UPLOAD_FOLDER"],
                    filename
                )
            )

            files.insert_one({
                "fullname": session["fullname"],
                "filename": filename,
                "filetype": filename.split(".")[-1].upper(),
                "uploaded_at": datetime.now().strftime("%d-%m-%Y %I:%M %p"),
                "favorite": False
                
            })

        return redirect("/files")

    search = request.args.get("search")

    if search:

        all_files = files.find({
            "fullname": session["fullname"],
            "filename":{
                "$regex":search,
                "$options":"i"
            }
        }).sort("favorite", -1)

    else:

        all_files = files.find({
            "fullname": session["fullname"]
        }).sort("favorite", -1)

    return render_template(
        "files.html",
        files=all_files
    )
@app.route("/favorite-file/<file_id>")
def favorite_file(file_id):

    if "fullname" not in session:
        return redirect("/login")

    file = files.find_one({
        "_id": ObjectId(file_id)
    })

    if file:

        files.update_one(
            {"_id": ObjectId(file_id)},
            {
                "$set": {
                    "favorite": not file.get("favorite", False)
                }
            }
        )

    return redirect("/files")
@app.route("/delete-file/<file_id>")
def delete_file(file_id):

    if "fullname" not in session:
        return redirect("/login")

    file = files.find_one({
        "_id": ObjectId(file_id)
    })

    if file:

        path = os.path.join(
            app.config["UPLOAD_FOLDER"],
            file["filename"]
        )

        if os.path.exists(path):
            os.remove(path)

        files.delete_one({
            "_id": ObjectId(file_id)
        })

    return redirect("/files")

@app.route("/download/<filename>")
def download_file(filename):

    return send_from_directory(
        app.config["UPLOAD_FOLDER"],
        filename,
        as_attachment=True
    )

@app.route("/view/<filename>")
def view_file(filename):

    if "fullname" not in session:
        return redirect("/login")

    return send_from_directory(
        app.config["UPLOAD_FOLDER"],
        filename
    )



from werkzeug.utils import secure_filename

@app.route("/profile", methods=["GET", "POST"])
def profile():

    if "fullname" not in session:
        return redirect("/login")

    user = users.find_one({
        "fullname": session["fullname"]
    })

    if request.method == "POST":

     
        phone = request.form["phone"]
        bio = request.form["bio"]


        image = user.get("profile_image", "")

        profile = request.files.get("profile")

        if profile and profile.filename != "":

            filename = secure_filename(profile.filename)

            upload_folder = os.path.join(
                "static",
                "images",
                "profiles"
            )

            os.makedirs(upload_folder, exist_ok=True)

            profile.save(
                os.path.join(upload_folder, filename)
            )

            image = filename

        users.update_one(

            {"email": session["email"]},

            {
                "$set":{

                    "phone": phone,
                    "bio": bio,
                    "profile_image": image

                }

            }

        )

        return redirect("/profile")

    user = users.find_one({
        "fullname": session["fullname"]
    })

    total_notes = notes.count_documents({
        "fullname": session["fullname"]
    })

    total_tasks = tasks.count_documents({
        "fullname": session["fullname"]
    })

    total_analysis = history.count_documents({
        "username": session["fullname"]
    })

    return render_template(

        "profile.html",

        user=user,

        total_notes=total_notes,

        total_tasks=total_tasks,

        total_analysis=total_analysis

    )

@app.route("/edit-profile", methods=["GET", "POST"])
def edit_profile():

    if "fullname" not in session:
        return redirect("/login")

    user = users.find_one({
        "fullname": session["fullname"]
    })

    if request.method == "POST":

        fullname = request.form["fullname"]
        phone = request.form["phone"]
        bio = request.form["bio"]

        users.update_one(

            {"fullname": session["fullname"]},

            {
                "$set":{

                    "fullname": fullname,
                    "phone": phone,
                    "bio": bio

                }

            }

        )

        session["fullname"] = fullname

        return redirect("/profile")

    return render_template(
        "edit_profile.html",
        user=user
    )

@app.route("/change-password", methods=["GET", "POST"])
def change_password():

    if "fullname" not in session:
        return redirect("/login")

    if request.method == "POST":

        current_password = request.form["current_password"]
        new_password = request.form["new_password"]
        confirm_password = request.form["confirm_password"]

        user = users.find_one({
            "fullname": session["fullname"]
        })

        if not check_password_hash(user["password"], current_password):

            return render_template(
                "change_password.html",
                error="Current Password is Incorrect!"
            )

        if new_password != confirm_password:

            return render_template(
                "change_password.html",
                error="Passwords do not match!"
            )

        users.update_one(

            {"fullname": session["fullname"]},

            {
                "$set":{

                    "password": generate_password_hash(new_password)

                }

            }

        )

        return render_template(
            "change_password.html",
            success="Password Changed Successfully!"
        )

    return render_template("change_password.html")

@app.route("/upload-profile-photo", methods=["POST"])
def upload_profile_photo():

    if "fullname" not in session:
        return redirect("/login")

    profile = request.files.get("profile")

    if profile and profile.filename != "":

        filename = secure_filename(profile.filename)

        upload_folder = os.path.join(
            "static",
            "images",
            "profiles"
        )

        os.makedirs(upload_folder, exist_ok=True)

        filepath = os.path.join(upload_folder, filename)

        profile.save(filepath)

        users.update_one(

            {"fullname": session["fullname"]},

            {
                "$set": {
                    "profile_image": filename
                }
            }

        )

    return redirect("/profile")

@app.route("/remove-profile-photo")
def remove_profile_photo():

    if "fullname" not in session:
        return redirect("/login")

    user = users.find_one({
        "fullname": session["fullname"]
    })

    if user and user.get("profile_image"):

        image_path = os.path.join(

            "static",
            "images",
            "profiles",

            user["profile_image"]

        )

        if os.path.exists(image_path):

            os.remove(image_path)

    users.update_one(

        {"fullname": session["fullname"]},

        {
            "$set": {
                "profile_image": ""
            }
        }

    )

    return redirect("/profile")

@app.route("/ai-analyzer")
def ai_analyzer():

    if "fullname" not in session:
        return redirect("/login")

    return render_template("ai_analyzer.html")


@app.route("/analyze-file", methods=["GET", "POST"])
def analyze_file():

    if "fullname" not in session:
        return redirect("/login")

    result = None
    filename = ""

    if request.method == "POST":

        file = request.files.get("file")

        if file and file.filename != "":

            filename = file.filename

  

            try:

                # Read Uploaded File
                file_text = file.read().decode("utf-8", errors="ignore")

                ext = os.path.splitext(filename)[1].lower()

                file.seek(0)

                file_text = ""


                if ext in [".txt", ".py", ".html", ".css", ".js", ".json", ".md"]:


                    file_text = file.read().decode("utf-8", errors="ignore")

                elif ext == ".pdf":
                    pdf = fitz.open(stream=file.read(), filetype="pdf")

                    for page in pdf:

                        file_text += page.get_text()


                elif ext == ".docx":
                    doc = Document(file)

                    for para in doc.paragraphs:

                        file_text += para.text + "\n"


                elif ext == ".pptx":
                    prs = Presentation(file)

                    for slide in prs.slides:
                        for shape in slide.shapes:
                            if hasattr(shape, "text"):
                                file_text += shape.text + "\n"



                elif ext in [".xlsx", ".xls"]:

                    df = pd.read_excel(file)

                    file_text = df.to_string(index=False)


                elif ext == ".csv":

                    df = pd.read_csv(file)
                    file_text = df.to_string(index=False)



                elif ext in [".png", ".jpg", ".jpeg", ".webp"]:
                    file.seek(0)
                   
                    image = PIL.Image.open(file)
                    

                    response = ai_client.models.generate_content(
                        model="gemini-3.5-flash",
                        contents=[
                            "Analyze this image professionally.",
                            image
                            
                        ]
                    )
                    result = response.text

                else:
                    file_text = "Unsupported File Format"

                prompt = f"""
You are an Expert AI File Analyzer.

Analyze the uploaded file professionally.

IMPORTANT RULES:

- Return only plain text.
- Do NOT return JSON.
- Do NOT return Markdown.
- Use exactly these headings.

📄 File Overview

📝 Summary

📌 Key Points

⭐ Keywords

❓ Important Questions

💡 Suggestions

📊 Overall Overview

🏆 Quality Score

Analyze this file:

{file_text}
"""

                try:

                    completion = groq_client.chat.completions.create(
                        model="llama-3.3-70b-versatile",
                        messages=[
                            {
                                "role": "user",
                                "content": prompt
                            }
                        ]
                    )

                    result = completion.choices[0].message.content
                    provider = "Groq"

                    save_history(
                        session["fullname"],
                        "File Analysis",
                        filename,
                        result,
                        "Groq"
                    )

                except Exception:

                    try:

                        response = ai_client.models.generate_content(
                            model="gemini-3.5-flash",
                            contents=prompt
                        )

                        result = response.text
                        provider = "Gemini"

                        save_history(
                            session["fullname"],
                            "File Analysis",
                            filename,
                            result,
                            "Gemini"
                        )

                    except Exception as e:

                        result = f"❌ Error : {str(e)}"


            except Exception as e:
                result = f"❌ Error : {str(e)}"  

    return render_template(
        "analyze_file.html",
        result=result,
        filename=filename
    )

@app.route("/analyze-project", methods=["GET", "POST"])
def analyze_project():

    if "fullname" not in session:
        return redirect("/login")

    result = None
    filename = ""

    if request.method == "POST":

        project = request.files.get("project")

        if project and project.filename != "":

            filename = project.filename

            
            try:

                upload_folder = "uploads"
                extract_folder = "uploads/project"

                os.makedirs(upload_folder, exist_ok=True)

                zip_path = os.path.join(upload_folder, project.filename)

                project.save(zip_path)

                def remove_readonly(func, path, exc_info):
                    os.chmod(path, stat.S_IWRITE)
                    func(path)


                if os.path.exists(extract_folder):
                    shutil.rmtree(extract_folder,
                onexc=remove_readonly)


                os.makedirs(extract_folder)

                with zipfile.ZipFile(zip_path, "r") as zip_ref:
                    zip_ref.extractall(extract_folder)

                project_text = ""

                

                for root, dirs, files in os.walk(extract_folder):
                    dirs[:] = [
                        d for d in dirs
                        if d not in [
                            "__pycache__",
                            ".git",
                            ".idea",
                            ".vscode",
                            "venv",
                            "env",
                            "node_modules"
                        ]
                    ]

                    

                    for file in files:

                        path = os.path.join(root, file)

                        if not os.path.isfile(path):
                            continue
                        

                        ext = os.path.splitext(file)[1].lower()

                        if ext in [".py", ".html", ".css", ".js", ".txt", ".md"]:
                            continue

                            try:

                                with open(path, "r", encoding="utf-8", errors="ignore") as f:

                                    project_text += f"\n\n===== {file} =====\n"

                                    project_text += f.read()

                            except Exception as e:
                                print("skipped:",path,e)
                                continue

                prompt = f"""
You are an Expert Software Engineer.

Analyze this complete software project.

IMPORTANT:

Return only plain text.

Use exactly these headings.

📁 Project Overview

📝 Project Summary

⚙ Technologies Used

⭐ Features Found

❌ Bugs Found

🔒 Security Analysis

⚡ Performance Analysis

💡 Suggestions

❓ Viva Questions

🏆 Overall Project Score

Project Files:

{project_text}
"""

                try:

                    completion = groq_client.chat.completions.create(
                        model="llama-3.3-70b-versatile",
                        messages=[
                            {
                                "role": "user",
                                "content": prompt
                            }
                        ]
                    )

                    result = completion.choices[0].message.content
                    provider = "Groq"

                    save_history(
                        session["fullname"],
                        "Project Analysis",
                        filename,
                        result,
                        "Groq"
                    )

                except Exception:
                    try:
                        response = ai_client.models.generate_content( 
                            model="gemini-3.5-flash",
                            contents=prompt

                        )

                        result = response.text
                        provider = "Gemini"

                        save_history(
                            session["fullname"],
                            "Project Analysis",
                            filename,
                            result,
                            "Gemini"
                        )

                    except Exception as e:

                            traceback.print_exc()

                            result = f"❌ Error: {str(e)}"


            except Exception as e:

                traceback.print_exc()

                result = f"❌ Error : {str(e)}"            

               
    return render_template(
        "analyze_project.html",
        result=result,
        filename=filename
    )


@app.route("/analyze-code", methods=["GET", "POST"])
def analyze_code():

    if "fullname" not in session:
        return redirect("/login")

    result = None
    code = ""

    if request.method == "POST":

        code = request.form["code"]

        prompt = f"""
You are an Expert Software Engineer.

Analyze the following code professionally.

IMPORTANT RULES:
- Do NOT return JSON.
- Do NOT return Python dictionary.
- Return only plain text.
- Use headings exactly as below.

Format:

🌐 Programming Language

📝 Summary

❌ Errors
- Error 1
- Error 2

💡 Suggestions
- Suggestion 1
- Suggestion 2

🔒 Security

⚡ Performance

🔧 Fixed Code

⭐ AI Score

Now analyze this code:

{code}
"""

        try:

            completion = groq_client.chat.completions.create(

                model="llama-3.3-70b-versatile",

                messages=[
                   {
                        "role": "user",
                        "content": prompt
                   }
                ]
            )

            result = completion.choices[0].message.content

            provider = "Groq"

            save_history(
                session["fullname"],
                "Code Analysis",
                "Code Analysis",
                result,
                "Groq"

            )

        except Exception:

            try:
            

                response = ai_client.models.generate_content(
                model="gemini-3.5-flash",
                contents=prompt

                )

                result = response.text
                provider = "Gemini"

                save_history(
                    session["fullname"],
                    "Code Analysis",
                    "Code Analysis",
                    result,
                    "Gemini"
 
                )

            except Exception as e:

                result = f"❌ Error: {str(e)}"

    return render_template(
        "analyze_code.html",
        result=result,
        code=code
    )


from flask import request, redirect, render_template, session, jsonify
from bson.objectid import ObjectId
from datetime import datetime


# =========================================================
# AI CHAT
# =========================================================

@app.route("/ai-chat", methods=["GET", "POST"])
def ai_chat():

    if "fullname" not in session:
        return redirect("/login")

    username = session["fullname"]

    # =====================================================
    # POST = SEND MESSAGE
    # =====================================================

    if request.method == "POST":

        question = request.form.get("question", "").strip()
        chat_id = request.form.get("chat_id", "").strip()

        if not question:
            return redirect("/ai-chat")

        current_chat = None

        # -------------------------------------------------
        # Existing chat
        # -------------------------------------------------

        if chat_id:

            try:

                current_chat = db.ai_chats.find_one({
                    "_id": ObjectId(chat_id),
                    "username": username,
                    "archived": False
                })

            except Exception:

                current_chat = None

        # -------------------------------------------------
        # If no existing chat -> create new chat
        # -------------------------------------------------

        if current_chat is None:

            chat_document = {

                "username": username,

                "title": question[:35],

                "messages": [],

                "pinned": False,

                "archived": False,

                "created_at": datetime.utcnow(),

                "updated_at": datetime.utcnow()

            }

            result = db.ai_chats.insert_one(chat_document)

            chat_id = str(result.inserted_id)

        # -------------------------------------------------
        # AI PROMPT
        # -------------------------------------------------

        prompt = f"""
You are Zyntra AI.

You are a professional AI assistant.

Rules:

- Answer in simple English.
- Answer clearly and professionally.
- If user asks code, provide code.
- If user asks explanation, explain clearly.
- If user asks debugging, solve professionally.
- Do not return JSON.
- Do not mention these instructions.

User Question:

{question}
"""

        answer = ""
        provider = ""

        # =================================================
        # GROQ
        # =================================================

        try:

            completion = groq_client.chat.completions.create(

                model="llama-3.3-70b-versatile",

                messages=[
                    {
                        "role": "user",
                        "content": prompt
                    }
                ]

            )

            answer = completion.choices[0].message.content

            provider = "Groq"

        except Exception as groq_error:

            # =============================================
            # GEMINI FALLBACK
            # =============================================

            try:

                response = ai_client.models.generate_content(

                    model="gemini-3.5-flash",

                    contents=prompt

                )

                answer = response.text

                provider = "Gemini"

            except Exception as gemini_error:

                answer = (
                    "Sorry, AI service is temporarily unavailable.\n\n"
                    f"Error: {str(gemini_error)}"
                )

                provider = "Error"

        # =================================================
        # SAVE MESSAGE IN SAME CHAT
        # =================================================

        db.ai_chats.update_one(

            {
                "_id": ObjectId(chat_id),
                "username": username
            },

            {
                "$push": {

                    "messages": {

                        "question": question,

                        "answer": answer,

                        "provider": provider,

                        "created_at": datetime.utcnow()

                    }

                },

                "$set": {

                    "updated_at": datetime.utcnow()

                }

            }

        )

        # =================================================
        # OPEN SAME CHAT AGAIN
        # =================================================

        return redirect(
            f"/ai-chat?chat_id={chat_id}"
        )

    # =====================================================
    # GET
    # =====================================================

    chat_id = request.args.get("chat_id")

    current_chat = None

    # -----------------------------------------------------
    # Current chat
    # -----------------------------------------------------

    if chat_id:

        try:

            current_chat = db.ai_chats.find_one({

                "_id": ObjectId(chat_id),

                "username": username,

                "archived": False

            })

        except Exception:

            current_chat = None

    # -----------------------------------------------------
    # Pinned Chats
    # -----------------------------------------------------

    pinned_chats = list(

        db.ai_chats.find({

            "username": username,

            "pinned": True,

            "archived": False

        }).sort(

            "updated_at",
            -1

        )

    )

    # -----------------------------------------------------
    # Recent Chats
    # -----------------------------------------------------

    recent_chats = list(

        db.ai_chats.find({

            "username": username,

            "pinned": False,

            "archived": False

        }).sort(

            "updated_at",
            -1

        )

    )

    # -----------------------------------------------------
    # Archived Chats
    # -----------------------------------------------------

    archived_chats = list(

        db.ai_chats.find({

            "username": username,

            "archived": True

        }).sort(

            "updated_at",
            -1

        )

    )

    return render_template(

        "ai_chat.html",

        current_chat=current_chat,

        pinned_chats=pinned_chats,

        recent_chats=recent_chats,

        archived_chats=archived_chats

    )

from flask import request, jsonify, session
from bson import ObjectId


# ============================================================
# DELETE CHAT
# ============================================================

@app.route("/ai-chat/delete/<chat_id>", methods=["POST"])
def delete_chat(chat_id):

    if "fullname" not in session:
        return jsonify({
            "success": False,
            "message": "Unauthorized"
        }), 401

    username = session["fullname"]

    try:

        result = db.ai_chats.delete_one({
            "_id": ObjectId(chat_id),
            "username": username
        })

        if result.deleted_count == 1:

            return jsonify({
                "success": True,
                "message": "Chat deleted successfully"
            })

        return jsonify({
            "success": False,
            "message": "Chat not found"
        }), 404

    except Exception as e:

        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


# ============================================================
# RENAME CHAT
# ============================================================

@app.route("/ai-chat/rename/<chat_id>", methods=["POST"])
def rename_chat(chat_id):

    if "fullname" not in session:
        return jsonify({
            "success": False,
            "message": "Unauthorized"
        }), 401

    username = session["fullname"]

    try:

        data = request.get_json(silent=True) or {}

        title = str(
            data.get("title", "")
        ).strip()

        if not title:

            return jsonify({
                "success": False,
                "message": "Chat name is required"
            }), 400

        if len(title) > 80:

            return jsonify({
                "success": False,
                "message": "Chat name is too long"
            }), 400

        result = db.ai_chats.update_one(
            {
                "_id": ObjectId(chat_id),
                "username": username
            },
            {
                "$set": {
                    "title": title
                }
            }
        )

        if result.matched_count == 0:

            return jsonify({
                "success": False,
                "message": "Chat not found"
            }), 404

        return jsonify({
            "success": True,
            "message": "Chat renamed successfully",
            "title": title
        })

    except Exception as e:

        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


# ============================================================
# PIN / UNPIN CHAT
# ============================================================

@app.route("/ai-chat/pin/<chat_id>", methods=["POST"])
def pin_chat(chat_id):

    if "fullname" not in session:
        return jsonify({
            "success": False,
            "message": "Unauthorized"
        }), 401

    username = session["fullname"]

    try:

        chat = db.ai_chats.find_one({
            "_id": ObjectId(chat_id),
            "username": username
        })

        if not chat:

            return jsonify({
                "success": False,
                "message": "Chat not found"
            }), 404

        new_value = not bool(
            chat.get("pinned", False)
        )

        db.ai_chats.update_one(
            {
                "_id": ObjectId(chat_id),
                "username": username
            },
            {
                "$set": {
                    "pinned": new_value
                }
            }
        )

        return jsonify({
            "success": True,
            "pinned": new_value
        })

    except Exception as e:

        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


# ============================================================
# ARCHIVE CHAT
# ============================================================

@app.route("/ai-chat/archive/<chat_id>", methods=["POST"])
def archive_chat(chat_id):

    if "fullname" not in session:
        return jsonify({
            "success": False,
            "message": "Unauthorized"
        }), 401

    username = session["fullname"]

    try:

        result = db.ai_chats.update_one(
            {
                "_id": ObjectId(chat_id),
                "username": username
            },
            {
                "$set": {
                    "archived": True
                }
            }
        )

        if result.matched_count == 0:

            return jsonify({
                "success": False,
                "message": "Chat not found"
            }), 404

        return jsonify({
            "success": True,
            "message": "Chat archived"
        })

    except Exception as e:

        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


# ============================================================
# UNARCHIVE CHAT
# ============================================================

@app.route("/ai-chat/unarchive/<chat_id>", methods=["POST"])
def unarchive_chat(chat_id):

    if "fullname" not in session:
        return jsonify({
            "success": False,
            "message": "Unauthorized"
        }), 401

    username = session["fullname"]

    try:

        result = db.ai_chats.update_one(
            {
                "_id": ObjectId(chat_id),
                "username": username
            },
            {
                "$set": {
                    "archived": False
                }
            }
        )

        if result.matched_count == 0:

            return jsonify({
                "success": False,
                "message": "Chat not found"
            }), 404

        return jsonify({
            "success": True,
            "message": "Chat unarchived"
        })

    except Exception as e:

        return jsonify({
            "success": False,
            "message": str(e)
        }), 500


@app.route("/settings")
def settings():

    if "fullname" not in session:
        return redirect("/login")

    user = users.find_one({
        "fullname": session["fullname"]
    })

    return render_template(
        "settings.html",
        user=user
    )
@app.route("/about")
def about():

    if "fullname" not in session:
        return redirect("/login")
    

    return render_template("about.html")

# ============================================================
# ZYNTRA AI FILE CREATOR - FINAL PROFESSIONAL VERSION
# ============================================================

import csv
import io
import json
import mimetypes
import re
import traceback
import uuid
import zipfile
from pathlib import Path
from datetime import datetime

from flask import (
    jsonify,
    request,
    render_template,
    redirect,
    session,
    send_file
)
from werkzeug.utils import secure_filename
# ============================================================
# FILE CREATOR PAGE
# ============================================================

@app.route("/file-creator", methods=["GET"])
def file_creator_page():
    try:
        if "fullname" not in session:
            return redirect("/login")

        username = session.get("fullname")

        recent_files = []

        try:
            recent_files = list(
                files.find(
                    {
                        "fullname": username,
                        "creator": "file_creator"
                    }
                )
                .sort("created_at", -1)
                .limit(6)
            )
        except Exception as e:
            print("File Creator MongoDB Error:", e)

        return render_template(
            "ai_file_creator.html",
            recent_files=recent_files
        )

    except Exception as e:
        print("FILE CREATOR PAGE ERROR:", e)
        traceback.print_exc()
        return "File Creator page error", 500

# ============================================================
# CONFIG
# ============================================================

GENERATED_DIR = Path("generated_files")
GENERATED_DIR.mkdir(parents=True, exist_ok=True)

FILE_EXTENSIONS = {
    "html": ".html",
    "css": ".css",
    "js": ".js",
    "javascript": ".js",
    "ts": ".ts",
    "typescript": ".ts",
    "jsx": ".jsx",
    "tsx": ".tsx",

    "python": ".py",
    "py": ".py",
    "c": ".c",
    "cpp": ".cpp",
    "c++": ".cpp",
    "java": ".java",
    "php": ".php",
    "sh": ".sh",
    "bash": ".sh",
    "bat": ".bat",

    "json": ".json",
    "xml": ".xml",
    "yaml": ".yaml",
    "yml": ".yml",
    "csv": ".csv",
    "sql": ".sql",

    "ini": ".ini",
    "conf": ".conf",
    "env": ".env",
    "log": ".log",

    "txt": ".txt",
    "md": ".md",
    "markdown": ".md",

    "pdf": ".pdf",
    "docx": ".docx",
    "xlsx": ".xlsx",
    "pptx": ".pptx",
    "zip": ".zip",
}


CODE_TYPES = {
    "html",
    "css",
    "js",
    "javascript",
    "ts",
    "typescript",
    "jsx",
    "tsx",
    "python",
    "py",
    "c",
    "cpp",
    "c++",
    "java",
    "php",
    "sh",
    "bash",
    "bat",
    "json",
    "xml",
    "yaml",
    "yml",
    "csv",
    "sql",
    "ini",
    "conf",
    "env",
    "log",
    "txt",
    "md",
    "markdown",
}


# ============================================================
# NORMALIZE TYPE
# ============================================================

def normalize_file_type(value):

    value = str(value or "txt").strip().lower()

    aliases = {
        "javascript": "js",
        "typescript": "ts",
        "python": "py",
        "c++": "cpp",
        "bash": "sh",
        "markdown": "md",
        "yml": "yaml",
    }

    return aliases.get(value, value)


# ============================================================
# SAFE FILE NAME
# ============================================================

def safe_filename(name, default="zyntra_file"):

    name = str(name or "").strip()

    name = Path(name).name

    name = secure_filename(name)

    if not name:
        name = default

    return name


def ensure_extension(filename, file_type):

    file_type = normalize_file_type(file_type)

    extension = FILE_EXTENSIONS.get(
        file_type,
        ".txt"
    )

    filename = safe_filename(filename)

    current_ext = Path(filename).suffix.lower()

    valid_extensions = {
        ext.lower()
        for ext in FILE_EXTENSIONS.values()
    }

    if current_ext in valid_extensions:
        filename = filename[:-len(current_ext)]

    if not filename:
        filename = "zyntra_file"

    return filename + extension


# ============================================================
# CLEAN AI RESPONSE
# ============================================================

def clean_ai_output(text):

    if text is None:
        return ""

    text = str(text).strip()

    text = re.sub(
        r"^\s*```[^\n]*\n",
        "",
        text
    )

    text = re.sub(
        r"\n?\s*```\s*$",
        "",
        text
    )

    return text.strip()


# ============================================================
# LOGIN
# ============================================================

def file_creator_logged_in():

    return bool(
        session.get("fullname")
    )


# ============================================================
# AUTO TYPE DETECTION
# ============================================================

def detect_file_type(prompt):

    text = str(prompt or "").lower()

    checks = [

        ("zip", [
            "project zip",
            "zip project",
            "zip file",
            "source code zip",
            "project files"
        ]),

        ("pptx", [
            "powerpoint",
            "presentation",
            "ppt",
            "pptx",
            "slide deck",
            "slides"
        ]),

        ("xlsx", [
            "excel",
            "spreadsheet",
            "workbook",
            "xlsx",
            "table"
        ]),

        ("docx", [
            "word document",
            "word file",
            "docx",
            "ms word"
        ]),

        ("pdf", [
            "pdf",
            "pdf document",
            "report pdf"
        ]),

        ("tsx", [
            "tsx",
            "react typescript"
        ]),

        ("jsx", [
            "jsx",
            "react component"
        ]),

        ("ts", [
            "typescript"
        ]),

        ("html", [
            "html",
            "web page",
            "website"
        ]),

        ("css", [
            "css",
            "stylesheet"
        ]),

        ("js", [
            "javascript",
            "node.js",
            "node js"
        ]),

        ("py", [
            "python",
            "python script"
        ]),

        ("cpp", [
            "c++",
            "cpp"
        ]),

        ("java", [
            "java program",
            "java code",
            "java file"
        ]),

        ("php", [
            "php"
        ]),

        ("sql", [
            "sql",
            "database query"
        ]),

        ("json", [
            "json"
        ]),

        ("xml", [
            "xml"
        ]),

        ("yaml", [
            "yaml",
            "yml"
        ]),

        ("csv", [
            "csv"
        ]),

        ("md", [
            "markdown"
        ]),
    ]

    for file_type, keywords in checks:

        for keyword in keywords:

            if keyword in text:
                return file_type

    return "txt"


# ============================================================
# AI SYSTEM PROMPT
# ============================================================

def build_system_prompt(file_type):

    file_type = normalize_file_type(
        file_type
    )

    common = """
You are Zyntra AI File Creator.

Your task is to generate the ACTUAL CONTENT that will be saved
inside the requested file.

STRICT RULES:

- Return ONLY final content.
- Never explain what you are doing.
- Never say "Here is your file".
- Never use Markdown code fences.
- Never provide instructions for creating the file.
- Never provide Python code for generating another file.
- Never return ReportLab code.
- Never return FPDF code.
- Never return openpyxl code.
- Never return python-docx code.
- Never return python-pptx code.
- Generate usable, complete content.
- Follow the user's requirements exactly.
"""

    rules = {

        "pdf": """
PDF:
Generate document CONTENT only.
Do not generate Python.
Do not generate PDF source syntax.
Do not generate %PDF-1.4.
Use professional headings, paragraphs, lists and sections.
""",

        "docx": """
DOCX:
Generate document CONTENT only.
Use headings, paragraphs, bullets and numbered sections.
Do not generate Python or document-generation code.
""",

        "xlsx": """
XLSX:
Generate DATA only.
Return CSV-style tabular data.
First row must normally be column headers.
Use comma-separated values.
Do not generate Python.
Do not explain the spreadsheet.
""",

        "csv": """
CSV:
Return valid CSV only.
First row should contain headers.
Use proper CSV quoting.
""",

        "pptx": """
PPTX:
Generate presentation CONTENT only.
Use clear slide sections.
Use headings beginning with # or ##.
Use bullet points beginning with -.
Do not generate PowerPoint Python code.
""",

        "zip": """
ZIP:
The response will be interpreted as a JSON project manifest.
Return ONLY valid JSON.
Format:

{
  "files": [
    {
      "path": "index.html",
      "content": "..."
    }
  ]
}

Use relative paths.
Never use ../.
Never use absolute paths.
Generate all required project files.
""",

        "html": """
HTML:
Return complete valid HTML5.
Include DOCTYPE, html, head and body.
""",

        "css": """
CSS:
Return CSS only.
No HTML.
No explanation.
""",

        "js": """
JavaScript:
Return valid JavaScript only.
""",

        "ts": """
TypeScript:
Return valid TypeScript only.
""",

        "jsx": """
JSX:
Return valid React JSX only.
""",

        "tsx": """
TSX:
Return valid React TypeScript only.
""",

        "py": """
Python:
Return valid runnable Python.
Use correct indentation.
""",

        "java": """
Java:
Return valid Java source code.
Use proper class structure.
""",

        "cpp": """
C++:
Return valid C++ source code.
""",

        "c": """
C:
Return valid C source code.
""",

        "php": """
PHP:
Return valid PHP source code.
""",

        "json": """
JSON:
Return STRICT valid JSON.
No comments.
No trailing commas.
""",

        "xml": """
XML:
Return valid XML.
""",

        "yaml": """
YAML:
Return valid YAML.
""",

        "sql": """
SQL:
Return valid SQL.
""",

        "md": """
Markdown:
Return clean Markdown.
""",

        "txt": """
Text:
Return clean readable plain text.
"""
    }

    return (
        common
        + "\n"
        + rules.get(file_type, "")
    )


# ============================================================
# GROQ
# ============================================================

def generate_with_groq(
    system_prompt,
    user_prompt
):

    if not GROQ_API_KEY:
        raise RuntimeError(
            "GROQ_API_KEY is not configured."
        )

    response = groq_client.chat.completions.create(

        model="llama-3.3-70b-versatile",

        messages=[
            {
                "role": "system",
                "content": system_prompt
            },
            {
                "role": "user",
                "content": user_prompt
            }
        ],

        temperature=0.15,

        max_tokens=32000
    )

    return clean_ai_output(
        response.choices[0]
        .message
        .content
    )


# ============================================================
# GEMINI FALLBACK
# ============================================================

def generate_with_gemini(
    system_prompt,
    user_prompt
):

    if not GEMINI_API_KEY:
        raise RuntimeError(
            "GEMINI_API_KEY is not configured."
        )

    response = ai_client.models.generate_content(

        model="gemini-3.5-flash",

        contents=[
            system_prompt,
            user_prompt
        ]
    )

    return clean_ai_output(
        response.text
    )


# ============================================================
# MARKDOWN PARSER
# ============================================================

def parse_blocks(content):

    blocks = []

    for raw in str(content or "").splitlines():

        line = raw.strip()

        if not line:
            blocks.append({
                "type": "space",
                "text": ""
            })
            continue

        match = re.match(
            r"^(#{1,6})\s+(.+)$",
            line
        )

        if match:

            level = len(
                match.group(1)
            )

            block_type = (
                "title"
                if level == 1
                else "heading"
                if level <= 2
                else "subheading"
            )

            blocks.append({
                "type": block_type,
                "text": match.group(2)
            })

            continue

        if re.match(
            r"^[-*•]\s+",
            line
        ):

            blocks.append({
                "type": "bullet",
                "text": re.sub(
                    r"^[-*•]\s+",
                    "",
                    line
                )
            })

            continue

        if re.match(
            r"^\d+[.)]\s+",
            line
        ):

            blocks.append({
                "type": "number",
                "text": re.sub(
                    r"^\d+[.)]\s+",
                    "",
                    line
                )
            })

            continue

        blocks.append({
            "type": "paragraph",
            "text": line
        })

    return blocks


def plain_text(text):

    text = str(text or "")

    text = re.sub(
        r"\*\*(.*?)\*\*",
        r"\1",
        text
    )

    text = re.sub(
        r"__(.*?)__",
        r"\1",
        text
    )

    text = re.sub(
        r"`([^`]*)`",
        r"\1",
        text
    )

    text = re.sub(
        r"\[(.*?)\]\(.*?\)",
        r"\1",
        text
    )

    return text.strip()


# ============================================================
# PROFESSIONAL PDF
# ============================================================

def create_pdf_file(
    path,
    content,
    title="Zyntra AI Document"
):

    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import (
        getSampleStyleSheet,
        ParagraphStyle
    )
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer
    )
    from xml.sax.saxutils import escape

    styles = getSampleStyleSheet()

    pdf_title = ParagraphStyle(
        "ZyntraPDFTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=27,
        alignment=TA_CENTER,
        spaceAfter=15
    )

    h1 = ParagraphStyle(
        "ZyntraH1",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=19,
        spaceBefore=12,
        spaceAfter=7
    )

    h2 = ParagraphStyle(
        "ZyntraH2",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        spaceBefore=9,
        spaceAfter=5
    )

    body = ParagraphStyle(
        "ZyntraBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=15,
        spaceAfter=7
    )

    bullet = ParagraphStyle(
        "ZyntraBullet",
        parent=body,
        leftIndent=15,
        firstLineIndent=-8,
        spaceAfter=5
    )

    number = ParagraphStyle(
        "ZyntraNumber",
        parent=body,
        leftIndent=18,
        firstLineIndent=-14,
        spaceAfter=5
    )

    doc = SimpleDocTemplate(

        str(path),

        pagesize=A4,

        rightMargin=20 * mm,
        leftMargin=20 * mm,

        topMargin=22 * mm,
        bottomMargin=20 * mm,

        title=title,
        author="Zyntra AI"
    )

    story = []

    blocks = parse_blocks(
        content
    )

    has_title = False

    for block in blocks:

        block_type = block["type"]

        text = plain_text(
            block["text"]
        )

        if block_type == "space":

            story.append(
                Spacer(1, 5)
            )

            continue

        safe = escape(text)

        if block_type == "title":

            story.append(
                Paragraph(
                    safe,
                    pdf_title
                )
            )

            has_title = True

        elif block_type == "heading":

            story.append(
                Paragraph(
                    safe,
                    h1
                )
            )

        elif block_type == "subheading":

            story.append(
                Paragraph(
                    safe,
                    h2
                )
            )

        elif block_type == "bullet":

            story.append(
                Paragraph(
                    f"• {safe}",
                    bullet
                )
            )

        elif block_type == "number":

            story.append(
                Paragraph(
                    safe,
                    number
                )
            )

        else:

            if not has_title:

                story.append(
                    Paragraph(
                        escape(title),
                        pdf_title
                    )
                )

                has_title = True

            story.append(
                Paragraph(
                    safe,
                    body
                )
            )

    if not story:

        story.append(
            Paragraph(
                escape(title),
                pdf_title
            )
        )

    def footer(canvas, document):

        canvas.saveState()

        width, height = A4

        canvas.setStrokeColor(
            colors.lightgrey
        )

        canvas.line(
            20 * mm,
            14 * mm,
            width - 20 * mm,
            14 * mm
        )

        canvas.setFont(
            "Helvetica",
            8
        )

        canvas.setFillColor(
            colors.grey
        )

        canvas.drawString(
            20 * mm,
            9 * mm,
            "Generated by Zyntra AI"
        )

        canvas.drawRightString(
            width - 20 * mm,
            9 * mm,
            f"Page {document.page}"
        )

        canvas.restoreState()

    doc.build(
        story,
        onFirstPage=footer,
        onLaterPages=footer
    )


# ============================================================
# PROFESSIONAL DOCX
# ============================================================

def create_docx_file(
    path,
    content,
    title="Zyntra AI Document"
):

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import (
        Inches,
        Pt
    )
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]

    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)

    # Header
    header = section.header

    header_p = header.paragraphs[0]

    header_p.text = "Zyntra AI"
    header_p.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    # Footer
    footer = section.footer

    footer_p = footer.paragraphs[0]

    footer_p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = footer_p.add_run(
        "Generated by Zyntra AI • Page "
    )

    fld = OxmlElement(
        "w:fldSimple"
    )

    fld.set(
        qn("w:instr"),
        "PAGE"
    )

    footer_p._p.append(fld)

    blocks = parse_blocks(
        content
    )

    title_done = False

    for block in blocks:

        block_type = block["type"]

        text = plain_text(
            block["text"]
        )

        if block_type == "space":
            continue

        if block_type == "title":

            p = document.add_paragraph()

            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            r = p.add_run(text)

            r.bold = True
            r.font.size = Pt(22)

            title_done = True

        elif block_type == "heading":

            p = document.add_heading(
                text,
                level=1
            )

            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(5)

        elif block_type == "subheading":

            document.add_heading(
                text,
                level=2
            )

        elif block_type == "bullet":

            document.add_paragraph(
                text,
                style="List Bullet"
            )

        elif block_type == "number":

            document.add_paragraph(
                text,
                style="List Number"
            )

        else:

            if not title_done:

                p = document.add_paragraph()

                p.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                )

                r = p.add_run(title)

                r.bold = True
                r.font.size = Pt(20)

                title_done = True

            p = document.add_paragraph(
                text
            )

            p.paragraph_format.space_after = Pt(7)
            p.paragraph_format.line_spacing = 1.15

    document.save(
        str(path)
    )


# ============================================================
# PROFESSIONAL XLSX
# ============================================================

def create_xlsx_file(
    path,
    content
):

    from openpyxl import Workbook
    from openpyxl.styles import (
        Font,
        PatternFill,
        Alignment,
        Border,
        Side
    )
    from openpyxl.utils import get_column_letter

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Zyntra Data"

    content = str(
        content or ""
    ).strip()

    rows = []

    # Try CSV
    try:

        reader = csv.reader(
            io.StringIO(content)
        )

        rows = [
            row
            for row in reader
            if row
        ]

    except Exception:

        rows = []

    # TSV fallback
    if (
        len(rows) > 0
        and len(rows[0]) == 1
        and "\t" in content
    ):

        reader = csv.reader(
            io.StringIO(content),
            delimiter="\t"
        )

        rows = [
            row
            for row in reader
            if row
        ]

    # Plain text fallback
    if not rows:

        rows = [
            [line]
            for line in content.splitlines()
            if line.strip()
        ]

    if not rows:

        rows = [
            ["Zyntra AI"],
            ["No data generated."]
        ]

    # Write
    for row_index, row in enumerate(
        rows,
        start=1
    ):

        for col_index, value in enumerate(
            row,
            start=1
        ):

            value = str(
                value or ""
            ).strip()

            cell = sheet.cell(
                row=row_index,
                column=col_index
            )

            # Preserve leading-zero strings
            if re.fullmatch(
                r"-?\d+",
                value
            ) and not (
                len(value.lstrip("-")) > 1
                and value.lstrip("-").startswith("0")
            ):

                try:
                    cell.value = int(value)
                except Exception:
                    cell.value = value

            elif re.fullmatch(
                r"-?\d+\.\d+",
                value
            ):

                try:
                    cell.value = float(value)
                except Exception:
                    cell.value = value

            elif value.lower() in {
                "true",
                "false"
            }:

                cell.value = (
                    value.lower() == "true"
                )

            else:

                cell.value = value

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    # Header
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="DCE6F1"
    )

    border_side = Side(
        style="thin",
        color="B7C3D0"
    )

    for cell in sheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.fill = header_fill

        cell.border = Border(
            bottom=border_side
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    # Freeze
    sheet.freeze_panes = "A2"

    # Filter
    if sheet.max_row >= 1:

        sheet.auto_filter.ref = (
            sheet.dimensions
        )

    # Width
    for column in range(
        1,
        sheet.max_column + 1
    ):

        max_length = 0

        for row in range(
            1,
            sheet.max_row + 1
        ):

            value = sheet.cell(
                row=row,
                column=column
            ).value

            if value is not None:

                max_length = max(
                    max_length,
                    len(str(value))
                )

        sheet.column_dimensions[
            get_column_letter(column)
        ].width = min(
            max(
                max_length + 2,
                12
            ),
            45
        )

    # Row height
    for row in range(
        1,
        sheet.max_row + 1
    ):

        sheet.row_dimensions[
            row
        ].height = 22

    workbook.save(
        str(path)
    )


# ============================================================
# PROFESSIONAL PPTX
# ============================================================

def create_pptx_file(
    path,
    content,
    title="Zyntra AI Presentation"
):

    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()

    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blocks = parse_blocks(
        content
    )

    presentation_title = title

    for block in blocks:

        if block["type"] == "title":

            presentation_title = (
                plain_text(
                    block["text"]
                )
            )

            break

    # --------------------------------------------------------
    # TITLE SLIDE
    # --------------------------------------------------------

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    title_box = slide.shapes.add_textbox(
        Inches(0.8),
        Inches(2.1),
        Inches(11.7),
        Inches(1.3)
    )

    tf = title_box.text_frame

    tf.clear()

    p = tf.paragraphs[0]

    p.text = presentation_title

    p.font.size = Pt(34)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    sub_box = slide.shapes.add_textbox(
        Inches(2),
        Inches(3.7),
        Inches(9.3),
        Inches(0.7)
    )

    sub_tf = sub_box.text_frame

    sub_tf.text = (
        "Generated by Zyntra AI"
    )

    sub_tf.paragraphs[0].font.size = Pt(16)
    sub_tf.paragraphs[0].alignment = (
        PP_ALIGN.CENTER
    )

    # --------------------------------------------------------
    # CONTENT SLIDES
    # --------------------------------------------------------

    current_title = None
    current_items = []

    def flush_slide():

        nonlocal current_title
        nonlocal current_items

        if not current_title and not current_items:
            return

        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        title_box = slide.shapes.add_textbox(
            Inches(0.7),
            Inches(0.45),
            Inches(11.8),
            Inches(0.75)
        )

        title_tf = title_box.text_frame

        title_tf.text = (
            plain_text(
                current_title
                or "Content"
            )
        )

        title_tf.paragraphs[0].font.size = Pt(27)
        title_tf.paragraphs[0].font.bold = True

        body_box = slide.shapes.add_textbox(
            Inches(0.9),
            Inches(1.4),
            Inches(11.4),
            Inches(5.4)
        )

        body_tf = body_box.text_frame

        body_tf.clear()
        body_tf.word_wrap = True

        for index, item in enumerate(
            current_items
        ):

            p = (
                body_tf.paragraphs[0]
                if index == 0
                else body_tf.add_paragraph()
            )

            text = plain_text(
                item["text"]
            )

            if item["type"] in {
                "bullet",
                "number"
            }:

                p.text = "• " + text

            else:

                p.text = text

            p.font.size = Pt(18)
            p.space_after = Pt(8)

        current_title = None
        current_items = []

    for block in blocks:

        if block["type"] in {
            "title",
            "heading",
            "subheading"
        }:

            flush_slide()

            current_title = (
                block["text"]
            )

        elif block["type"] != "space":

            current_items.append(
                block
            )

            if len(current_items) >= 7:

                flush_slide()

    flush_slide()

    if len(prs.slides) == 1:

        slide = prs.slides.add_slide(
            prs.slide_layouts[6]
        )

        box = slide.shapes.add_textbox(
            Inches(1),
            Inches(2),
            Inches(11),
            Inches(2)
        )

        box.text_frame.text = (
            "Zyntra AI Generated Presentation"
        )

        box.text_frame.paragraphs[0].font.size = Pt(28)

    prs.save(
        str(path)
    )


# ============================================================
# ZIP JSON PARSER
# ============================================================

def parse_zip_manifest(content):

    content = clean_ai_output(
        content
    )

    try:

        data = json.loads(
            content
        )

    except Exception:

        match = re.search(
            r"\{.*\}",
            content,
            re.DOTALL
        )

        if not match:

            raise ValueError(
                "AI returned invalid ZIP project."
            )

        data = json.loads(
            match.group(0)
        )

    if not isinstance(
        data,
        dict
    ):

        raise ValueError(
            "Invalid ZIP manifest."
        )

    project_files = data.get(
        "files"
    )

    if not isinstance(
        project_files,
        list
    ):

        raise ValueError(
            "ZIP manifest has no files."
        )

    return project_files


# ============================================================
# ZIP BUILDER
# ============================================================

def create_zip_file(
    path,
    content
):

    project_files = parse_zip_manifest(
        content
    )

    if not project_files:

        raise ValueError(
            "ZIP project is empty."
        )

    if len(project_files) > 50:

        raise ValueError(
            "ZIP project contains too many files."
        )

    total_size = 0
    added = 0

    with zipfile.ZipFile(
        path,
        "w",
        zipfile.ZIP_DEFLATED
    ) as archive:

        for item in project_files:

            if not isinstance(
                item,
                dict
            ):
                continue

            relative_path = str(
                item.get(
                    "path",
                    ""
                )
            ).strip()

            file_content = str(
                item.get(
                    "content",
                    ""
                )
            )

            if not relative_path:
                continue

            relative_path = (
                relative_path
                .replace("\\", "/")
                .lstrip("/")
            )

            parts = Path(
                relative_path
            ).parts

            if (
                ".." in parts
                or ":" in relative_path
            ):

                raise ValueError(
                    f"Unsafe ZIP path: {relative_path}"
                )

            if len(relative_path) > 200:

                raise ValueError(
                    "ZIP filename too long."
                )

            encoded = file_content.encode(
                "utf-8"
            )

            if len(encoded) > 500_000:

                raise ValueError(
                    f"File too large: {relative_path}"
                )

            total_size += len(encoded)

            if total_size > 5_000_000:

                raise ValueError(
                    "ZIP project is too large."
                )

            archive.writestr(
                relative_path,
                encoded
            )

            added += 1

    if added == 0:

        raise ValueError(
            "No valid files in ZIP."
        )


# ============================================================
# BUILD FILE
# ============================================================

def build_generated_file(
    file_type,
    filename,
    content
):

    file_type = normalize_file_type(
        file_type
    )

    if file_type not in FILE_EXTENSIONS:

        file_type = "txt"

    final_filename = ensure_extension(
        filename,
        file_type
    )

    stored_filename = (
        f"{uuid.uuid4().hex[:12]}_"
        f"{final_filename}"
    )

    output_path = (
        GENERATED_DIR /
        stored_filename
    )

    if file_type in CODE_TYPES:

        create_text_file(
            output_path,
            content
        )

    elif file_type == "pdf":

        create_pdf_file(
            output_path,
            content,
            Path(final_filename).stem
        )

    elif file_type == "docx":

        create_docx_file(
            output_path,
            content,
            Path(final_filename).stem
        )

    elif file_type == "xlsx":

        create_xlsx_file(
            output_path,
            content
        )

    elif file_type == "pptx":

        create_pptx_file(
            output_path,
            content,
            Path(final_filename).stem
        )

    elif file_type == "zip":

        create_zip_file(
            output_path,
            content
        )

    else:

        create_text_file(
            output_path,
            content
        )

    return (
        stored_filename,
        final_filename,
        output_path
    )


# ============================================================
# TEXT
# ============================================================

def create_text_file(
    path,
    content
):

    path.write_text(
        str(content or ""),
        encoding="utf-8"
    )


# ============================================================
# API
# ============================================================

@app.route(
    "/api/file-creator",
    methods=["POST"]
)
def file_creator_api():

    if not file_creator_logged_in():

        return jsonify({
            "success": False,
            "error": "Please login first."
        }), 401

    try:

        data = request.get_json(
            silent=True
        ) or {}

        prompt = str(
            data.get(
                "prompt",
                ""
            )
        ).strip()

        if not prompt:

            return jsonify({
                "success": False,
                "error":
                    "Please enter what you want to create."
            }), 400

        if len(prompt) > 20000:

            return jsonify({
                "success": False,
                "error":
                    "Prompt is too long."
            }), 400

        # ----------------------------------------------------
        # TYPE
        # ----------------------------------------------------

        file_type = str(
            data.get(
                "file_type",
                "auto"
            )
        ).strip().lower()

        if file_type == "auto":

            file_type = detect_file_type(
                prompt
            )

        file_type = normalize_file_type(
            file_type
        )

        if file_type not in FILE_EXTENSIONS:

            return jsonify({
                "success": False,
                "error":
                    f"Unsupported file type: {file_type}"
            }), 400

        # ----------------------------------------------------
        # NAME
        # ----------------------------------------------------

        requested_name = safe_filename(
            data.get(
                "file_name",
                "zyntra_file"
            ),
            "zyntra_file"
        )

        # ----------------------------------------------------
        # AI PROMPT
        # ----------------------------------------------------

        if file_type == "zip":

            system_prompt = build_system_prompt(
                "zip"
            )

            user_prompt = build_zip_user_prompt(
                prompt,
                requested_name
            )

        else:

            system_prompt = build_system_prompt(
                file_type
            )

            user_prompt = f"""
Create the actual {file_type} content requested below.

USER REQUEST:
{prompt}

FILE NAME:
{requested_name}

PROFESSIONAL LAYOUT:
{bool(data.get("professional_layout", True))}

AI OPTIMIZATION:
{bool(data.get("ai_optimize", True))}

Return ONLY the final content.
"""

        # ----------------------------------------------------
        # GROQ
        # ----------------------------------------------------

        provider = "Groq"

        try:

            generated_content = (
                generate_with_groq(
                    system_prompt,
                    user_prompt
                )
            )

        except Exception as groq_error:

            print(
                "Groq error:",
                groq_error
            )

            # ------------------------------------------------
            # GEMINI
            # ------------------------------------------------

            provider = "Gemini"

            generated_content = (
                generate_with_gemini(
                    system_prompt,
                    user_prompt
                )
            )

        if not generated_content:

            raise RuntimeError(
                "AI returned empty content."
            )

        # ----------------------------------------------------
        # CREATE REAL FILE
        # ----------------------------------------------------

        (
            stored_filename,
            final_filename,
            output_path
        ) = build_generated_file(
            file_type,
            requested_name,
            generated_content
        )

        # ----------------------------------------------------
        # VERIFY
        # ----------------------------------------------------

        if (
            not output_path.exists()
            or not output_path.is_file()
        ):

            raise RuntimeError(
                "File creation failed."
            )

        file_size = output_path.stat().st_size

        if file_size <= 0:

            raise RuntimeError(
                "Generated file is empty."
            )

        # ----------------------------------------------------
        # DATABASE
        # ----------------------------------------------------

        username = session["fullname"]

        try:

            files.insert_one({

                "fullname": username,

                "username": username,

                "creator": "file_creator",

                "filename": final_filename,

                "stored_filename":
                    stored_filename,

                "file_type": file_type,

                "provider": provider,

                "prompt": prompt,

                "path": str(output_path),

                "size": file_size,

                "created_at":
                    datetime.now()
            })

        except Exception as db_error:

            print(
                "File Creator DB error:",
                db_error
            )

        # ----------------------------------------------------
        # HISTORY
        # ----------------------------------------------------

        try:

            save_history(
                username,
                "file_creator",
                final_filename,
                generated_content,
                provider
            )

        except Exception as history_error:

            print(
                "History error:",
                history_error
            )

        # ----------------------------------------------------
        # RESPONSE
        # ----------------------------------------------------

        return jsonify({

            "success": True,

            "provider": provider,

            "file_type": file_type,

            "filename": final_filename,

            "stored_filename":
                stored_filename,

            "size": file_size,

            "download_url":
                f"/download-generated/"
                f"{stored_filename}",

            "preview_url":
                f"/preview-generated/"
                f"{stored_filename}"
        })

    except Exception as error:

        traceback.print_exc()

        return jsonify({

            "success": False,

            "error": str(error)

        }), 500


# ============================================================
# ZIP USER PROMPT
# ============================================================

def build_zip_user_prompt(
    prompt,
    project_name
):

    return f"""
Create a complete working software project based on this request:

{prompt}

Project name:
{project_name}

Return ONLY valid JSON.

Required structure:

{{
  "files": [
    {{
      "path": "index.html",
      "content": "complete content"
    }},
    {{
      "path": "style.css",
      "content": "complete content"
    }},
    {{
      "path": "script.js",
      "content": "complete content"
    }}
  ]
}}

Rules:

- All files must be complete.
- Files must work together.
- Use correct relative paths.
- Never use ../
- Never use absolute paths.
- Do not include explanations.
- Do not include Markdown fences.
- Do not include Python code for generating the ZIP.
- Generate the actual project files.
"""


# ============================================================
# OWNERSHIP
# ============================================================

def user_owns_generated_file(
    stored_filename
):

    if not file_creator_logged_in():
        return False

    username = session["fullname"]

    try:

        record = files.find_one({

            "fullname": username,

            "creator": "file_creator",

            "stored_filename":
                stored_filename
        })

        return record is not None

    except Exception as error:

        print(
            "Ownership error:",
            error
        )

        return False


# ============================================================
# DOWNLOAD
# ============================================================

@app.route(
    "/download-generated/<filename>"
)
def download_generated(filename):

    if not file_creator_logged_in():

        return redirect("/login")

    stored_filename = Path(
        filename
    ).name

    if not user_owns_generated_file(
        stored_filename
    ):

        return "File not found.", 404

    path = (
        GENERATED_DIR /
        stored_filename
    )

    if (
        not path.exists()
        or not path.is_file()
    ):

        return "File not found.", 404

    original_name = stored_filename

    if "_" in stored_filename:

        original_name = stored_filename.split(
            "_",
            1
        )[1]

    mimetype = (
        mimetypes.guess_type(
            str(path)
        )[0]
        or "application/octet-stream"
    )

    return send_file(

        path,

        as_attachment=True,

        download_name=original_name,

        mimetype=mimetype
    )


# ============================================================
# PREVIEW
# ============================================================

@app.route(
    "/preview-generated/<filename>"
)
def preview_generated(filename):

    if not file_creator_logged_in():

        return redirect("/login")

    stored_filename = Path(
        filename
    ).name

    if not user_owns_generated_file(
        stored_filename
    ):

        return "File not found.", 404

    path = (
        GENERATED_DIR /
        stored_filename
    )

    if (
        not path.exists()
        or not path.is_file()
    ):

        return "File not found.", 404

    mimetype = (
        mimetypes.guess_type(
            str(path)
        )[0]
        or "application/octet-stream"
    )

    # IMPORTANT:
    # PDF browser ma actual PDF viewer ma open thase.
    # Raw %PDF source nahi dekhay.

    return send_file(

        path,

        as_attachment=False,

        mimetype=mimetype
    )

# ============================================================
# APPLICATION START
# ============================================================




if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=False)
